"""
exporters.py
------------
Generate Excel (point table + line segment table) and DXF from analysis JSON.
Returns file bytes so the API endpoint can send them directly.
"""

from __future__ import annotations
import io
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
import ezdxf


def to_excel(analysis: dict) -> bytes:
    """
    Build an Excel workbook with two sheets:
      Sheet 1 - Point Table:        Point | X | Y
      Sheet 2 - Line Segment Table: Segment | End 1 | End 2
    Returns bytes of the .xlsx file.
    """
    wb = openpyxl.Workbook()

    # ── Sheet 1: Points ──────────────────────────────────────────────────────
    ws_points = wb.active
    ws_points.title = "Point Table"

    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(color="FFFFFF", bold=True)

    headers_p = ["Point", f"X ({analysis.get('unit', '')})", f"Y ({analysis.get('unit', '')})"]
    for col, h in enumerate(headers_p, 1):
        cell = ws_points.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    for row, corner in enumerate(analysis["corners"], 2):
        ws_points.cell(row=row, column=1, value=corner["label"])
        ws_points.cell(row=row, column=2, value=corner["x"])
        ws_points.cell(row=row, column=3, value=corner["y"])

    ws_points.column_dimensions["A"].width = 10
    ws_points.column_dimensions["B"].width = 14
    ws_points.column_dimensions["C"].width = 14

    # ── Sheet 2: Line Segments ───────────────────────────────────────────────
    ws_segs = wb.create_sheet("Line Segment Table")

    headers_s = ["Line Segment", "End 1", "End 2"]
    for col, h in enumerate(headers_s, 1):
        cell = ws_segs.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    for row, seg in enumerate(analysis["line_segments"], 2):
        ws_segs.cell(row=row, column=1, value=seg["number"])
        ws_segs.cell(row=row, column=2, value=seg["end1"])
        ws_segs.cell(row=row, column=3, value=seg["end2"])

    ws_segs.column_dimensions["A"].width = 16
    ws_segs.column_dimensions["B"].width = 10
    ws_segs.column_dimensions["C"].width = 10

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def to_dxf(analysis: dict) -> bytes:
    """
    Build a DXF file from corners and line segments.
    Returns bytes of the .dxf file.
    """
    doc = ezdxf.new("R2010")
    msp = doc.modelspace()

    corner_map = {c["label"]: (float(c["x"]), float(c["y"])) for c in analysis["corners"]}

    for seg in analysis["line_segments"]:
        start = corner_map.get(seg["end1"])
        end = corner_map.get(seg["end2"])
        if start and end:
            msp.add_line(start, end)

    # Add point labels as text
    for corner in analysis["corners"]:
        msp.add_text(
            corner["label"],
            dxfattribs={
                "insert": (float(corner["x"]), float(corner["y"])),
                "height": 0.5,
            },
        )

    buf = io.StringIO()
    doc.write(buf)
    return buf.getvalue().encode("utf-8")